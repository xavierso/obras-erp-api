import io
import os
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import select, or_
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage

from app.core.deps import require_director, get_current_user, get_empresa_id
from app.database import get_db
from app.models.obra import Obra
from app.models.presupuesto import Presupuesto, CapituloPresupuesto, PartidaPresupuesto
from app.models.incidencia import Incidencia
from app.models.tarea import Tarea
from app.models.usuario import Usuario
from app.models.invitacion import Invitacion, EstadoInvitacion
from app.models.empresa import Empresa

router = APIRouter(prefix="/exportar", tags=["Exportación"])

def draw_excel_dashboard_header(ws, title: str, perfil: Empresa, kpi_data: dict):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    
    BG_MAIN = "1A202C"      # Fondo muy oscuro
    BG_CARD = "2D3748"      # Gris oscuro para tarjetas
    ACCENT = "0D9488"       # Teal / Cyan para el acento
    TEXT_MAIN = "F8FAFC"
    TEXT_MUTED = "94A3B8"
    
    font_family = "Segoe UI"
    
    fill_bg = PatternFill(start_color=BG_MAIN, end_color=BG_MAIN, fill_type="solid")
    fill_card = PatternFill(start_color=BG_CARD, end_color=BG_CARD, fill_type="solid")
    
    ws.sheet_view.showGridLines = False
    
    # Fondo global
    for r in range(1, 150):
        for c in range(1, 15):
            ws.cell(row=r, column=c).fill = fill_bg

    # --- ROW 1: Título Principal ---
    ws.cell(row=1, column=2, value="DIAM GESTIÓN").font = Font(name=font_family, size=22, bold=True, color=TEXT_MAIN)
    ws.cell(row=1, column=4, value="DE OBRAS").font = Font(name=font_family, size=22, bold=True, color=ACCENT)
    
    # --- ROW 3: Logo / Proyecto y Tabs ---
    ws.cell(row=3, column=2, value="Empresa").font = Font(name=font_family, size=9, color=TEXT_MUTED)
    ws.cell(row=4, column=2, value=perfil.nombre if perfil else "DIAM ERP").font = Font(name=font_family, size=11, color=TEXT_MAIN)
    
    tabs = ["Presupuestos", "Obras", "Incidencias", "Tareas", "Personal"]
    col_idx = 4
    for tab in tabs:
        c = ws.cell(row=3, column=col_idx, value=tab.upper())
        c.font = Font(name=font_family, size=11, color=TEXT_MAIN if tab == title else TEXT_MUTED)
        if tab == title:
            c.fill = PatternFill(start_color=ACCENT, end_color=ACCENT, fill_type="solid")
            c.font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=3, start_column=col_idx, end_row=4, end_column=col_idx)
        col_idx += 1
        
    # --- ROW 6-8: Tarjetas KPI ---
    # Box 1: Obras
    for r in range(6, 9):
        for c in range(2, 4):
            ws.cell(row=r, column=c).fill = fill_card
    ws.cell(row=6, column=2, value="OBRAS ACTIVAS").font = Font(name=font_family, size=10, color=TEXT_MUTED)
    ws.cell(row=7, column=2, value=kpi_data.get('obras_activas', 0)).font = Font(name=font_family, size=20, bold=True, color=ACCENT)
    ws.cell(row=7, column=2).alignment = Alignment(vertical="center")
    
    # Box 2: Presupuesto
    for r in range(6, 9):
        for c in range(4, 6):
            ws.cell(row=r, column=c).fill = fill_card
    ws.cell(row=6, column=4, value="PRESUPUESTADO TOTAL").font = Font(name=font_family, size=10, color=TEXT_MUTED)
    
    val_total = kpi_data.get('presupuesto_total', 0)
    c_monto = ws.cell(row=7, column=4, value=val_total)
    c_monto.font = Font(name=font_family, size=20, bold=True, color=TEXT_MAIN)
    c_monto.number_format = '#,##0.00 €'
    c_monto.alignment = Alignment(vertical="center")
    
    # Box 3: Incidencias
    for r in range(6, 9):
        for c in range(6, 8):
            ws.cell(row=r, column=c).fill = fill_card
    ws.cell(row=6, column=6, value="INCIDENCIAS ABIERTAS").font = Font(name=font_family, size=10, color=TEXT_MUTED)
    ws.cell(row=7, column=6, value=kpi_data.get('incidencias_abiertas', 0)).font = Font(name=font_family, size=18, bold=True, color="EF4444")
    ws.cell(row=7, column=6).alignment = Alignment(vertical="center")
    
    # Box 4: Equipo
    for r in range(6, 9):
        for c in range(8, 10):
            ws.cell(row=r, column=c).fill = fill_card
    ws.cell(row=6, column=8, value="EQUIPO EN CAMPO").font = Font(name=font_family, size=10, color=TEXT_MUTED)
    ws.cell(row=7, column=8, value=kpi_data.get('equipo', 0)).font = Font(name=font_family, size=18, bold=True, color="10B981")
    ws.cell(row=7, column=8).alignment = Alignment(vertical="center")
    
    # --- ROW 10: Barra de Título de la tabla ---
    ws.cell(row=10, column=2, value=f"RESUMEN DE {title.upper()}").font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    teal_fill = PatternFill(start_color=ACCENT, end_color=ACCENT, fill_type="solid")
    for c in range(2, 10):
        ws.cell(row=10, column=c).fill = teal_fill

    return 11

def apply_table_header(cell):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    cell.fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid") # Teal header
    cell.font = Font(color="FFFFFF", bold=True, name="Segoe UI", size=11)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.border = Border(bottom=Side(style="medium", color="0F766E"))

def apply_badge(cell, estado: str):
    from openpyxl.styles import Font, PatternFill, Alignment
    estado = (estado or "").upper()
    font_family = "Segoe UI"
    
    if estado in ["EN_EJECUCION", "ACTIVO"]:
        cell.fill = PatternFill(start_color="0EA5E9", end_color="0EA5E9", fill_type="solid") 
        cell.font = Font(name=font_family, color="FFFFFF", bold=True, size=9)
    elif estado in ["PENDIENTE", "PENDIENTE_APROBACION", "BORRADOR", "EN_ESTUDIO", "ENVIADO"]:
        cell.fill = PatternFill(start_color="EAB308", end_color="EAB308", fill_type="solid") 
        cell.font = Font(name=font_family, color="000000", bold=True, size=9)
    elif estado in ["INCIDENCIA", "BLOQUEADO", "CANCELADO"]:
        cell.fill = PatternFill(start_color="EF4444", end_color="EF4444", fill_type="solid") 
        cell.font = Font(name=font_family, color="FFFFFF", bold=True, size=9)
    elif estado in ["COMPLETADO", "FINALIZADO", "APROBADO", "RESUELTA"]:
        cell.fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid") 
        cell.font = Font(name=font_family, color="FFFFFF", bold=True, size=9)
    else:
        cell.fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid") 
        cell.font = Font(name=font_family, color="FFFFFF", bold=True, size=9)
        
    cell.alignment = Alignment(horizontal="center", vertical="center")

@router.get("/excel")
async def export_to_excel(
    admin: Usuario = Depends(require_director),
    db: AsyncSession = Depends(get_db),
    empresa_id: int = Depends(get_empresa_id)
):
    result = await db.execute(select(Empresa).where(Empresa.id == empresa_id))
    perfil = result.scalar_one_or_none()

    # Pre-calculate KPIs
    res_obras = await db.execute(select(Obra).where(Obra.empresa_id == empresa_id))
    obras = res_obras.scalars().all()
    obras_activas = sum(1 for o in obras if o.estado and o.estado.value not in ["cancelado", "finalizado"])
    
    res_pres = await db.execute(
        select(Presupuesto).options(
            selectinload(Presupuesto.capitulos)
            .selectinload(CapituloPresupuesto.partidas)
            .selectinload(PartidaPresupuesto.lineas_medicion)
        ).where(Presupuesto.empresa_id == empresa_id)
    )
    presupuestos = res_pres.scalars().all()
    presupuesto_total = sum(p.total or 0 for p in presupuestos)
    
    res_inc = await db.execute(select(Incidencia).options(selectinload(Incidencia.obra)).where(Incidencia.empresa_id == empresa_id))
    incidencias = res_inc.scalars().all()
    incidencias_abiertas = sum(1 for i in incidencias if i.estado and i.estado.value not in ["resuelta"])
    
    res_usu = await db.execute(select(Usuario).where(Usuario.empresa_id == empresa_id))
    equipo_count = len(res_usu.scalars().all())
    
    kpi_data = {
        "obras_activas": obras_activas,
        "presupuesto_total": presupuesto_total,
        "incidencias_abiertas": incidencias_abiertas,
        "equipo": equipo_count
    }

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    BG_MAIN = "1A202C"
    BG_CARD = "2D3748"
    TEXT_MAIN = "F8FAFC"
    
    def style_row(ws, row_idx, max_col, is_even):
        bg = BG_CARD if is_even else "232B38"
        from openpyxl.styles import PatternFill, Font, Border, Side
        fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
        for c in range(2, max_col + 1):
            cell = ws.cell(row=row_idx, column=c)
            
            # Identify if it's a badge by checking if it already has a non-default fill
            is_badge = False
            if cell.fill and cell.fill.start_color and type(cell.fill.start_color.rgb) == str:
                if cell.fill.start_color.rgb not in ["00000000", bg, "232B38", "2D3748", BG_MAIN, BG_CARD]:
                    is_badge = True
            
            if not is_badge:
                cell.fill = fill
                # Apply explicit colors based on column type
                # The code column is column 2 usually, let's just check the existing font color
                if cell.font and cell.font.color and getattr(cell.font.color, 'rgb', None) in ["FF94A3B8", "0094A3B8", "94A3B8"]:
                    cell.font = Font(name="Consolas", size=10, color="94A3B8")
                else:
                    # Force white text for standard cells
                    cell.font = Font(name="Segoe UI", size=10, color="FFFFFF")
                    
            cell.border = Border(bottom=Side(style="thin", color="1A202C"), right=Side(style="thin", color="1A202C"))

    # 1. Presupuestos
    ws_pres = wb.create_sheet(title="Presupuestos")
    start_row = draw_excel_dashboard_header(ws_pres, "Presupuestos", perfil, kpi_data)
    
    headers = ["CÓDIGO", "NOMBRE", "ESTADO", "MONTO"]
    for col, h in enumerate(headers, 2):
        cell = ws_pres.cell(row=start_row, column=col, value=h)
        apply_table_header(cell)
        
    for i, p in enumerate(presupuestos, start=start_row + 1):
        ws_pres.cell(row=i, column=2, value=p.codigo or "").font = Font(name="Consolas", color="94A3B8")
        ws_pres.cell(row=i, column=3, value=p.nombre)
        
        c_estado = ws_pres.cell(row=i, column=4, value=p.estado.value.upper() if p.estado else "")
        apply_badge(c_estado, p.estado.value if p.estado else "")
        
        c_monto = ws_pres.cell(row=i, column=5, value=p.total)
        c_monto.number_format = '#,##0.00 €'
        c_monto.alignment = Alignment(horizontal="right")
        
        style_row(ws_pres, i, 5, is_even=(i%2==0))
        
    # 2. Obras
    ws_obras = wb.create_sheet(title="Obras")
    start_row = draw_excel_dashboard_header(ws_obras, "Obras", perfil, kpi_data)
    
    headers_obras = ["CÓDIGO", "NOMBRE", "CLIENTE", "ESTADO"]
    for col, h in enumerate(headers_obras, 2):
        cell = ws_obras.cell(row=start_row, column=col, value=h)
        apply_table_header(cell)
        
    for i, o in enumerate(obras, start=start_row + 1):
        ws_obras.cell(row=i, column=2, value=o.codigo).font = Font(name="Consolas", color="94A3B8")
        ws_obras.cell(row=i, column=3, value=o.nombre)
        ws_obras.cell(row=i, column=4, value=o.cliente or "")
        c_estado = ws_obras.cell(row=i, column=5, value=o.estado.value.upper() if o.estado else "")
        apply_badge(c_estado, o.estado.value if o.estado else "")
        style_row(ws_obras, i, 5, is_even=(i%2==0))

    # 3. Incidencias
    ws_inc = wb.create_sheet(title="Incidencias")
    start_row = draw_excel_dashboard_header(ws_inc, "Incidencias", perfil, kpi_data)
    
    headers_inc = ["CÓDIGO", "TÍTULO", "OBRA", "ESTADO"]
    for col, h in enumerate(headers_inc, 2):
        cell = ws_inc.cell(row=start_row, column=col, value=h)
        apply_table_header(cell)
        
    for i, inc in enumerate(incidencias, start=start_row + 1):
        ws_inc.cell(row=i, column=2, value=inc.codigo).font = Font(name="Consolas", color="94A3B8")
        ws_inc.cell(row=i, column=3, value=inc.titulo)
        ws_inc.cell(row=i, column=4, value=inc.obra.nombre if inc.obra else "")
        c_estado = ws_inc.cell(row=i, column=5, value=inc.estado.value.upper() if inc.estado else "")
        apply_badge(c_estado, inc.estado.value if inc.estado else "")
        style_row(ws_inc, i, 5, is_even=(i%2==0))

    # 4. Tareas
    ws_tar = wb.create_sheet(title="Tareas")
    start_row = draw_excel_dashboard_header(ws_tar, "Tareas", perfil, kpi_data)
    
    res_tar = await db.execute(select(Tarea).options(selectinload(Tarea.obra)).where(Tarea.empresa_id == empresa_id))
    tareas = res_tar.scalars().all()
    
    headers_tar = ["TÍTULO", "OBRA", "FECHA LÍMITE", "ESTADO"]
    for col, h in enumerate(headers_tar, 2):
        cell = ws_tar.cell(row=start_row, column=col, value=h)
        apply_table_header(cell)
        
    for i, tar in enumerate(tareas, start=start_row + 1):
        ws_tar.cell(row=i, column=2, value=tar.titulo)
        ws_tar.cell(row=i, column=3, value=tar.obra.nombre if tar.obra else "")
        ws_tar.cell(row=i, column=4, value=tar.fecha_limite.strftime("%Y-%m-%d") if tar.fecha_limite else "")
        c_estado = ws_tar.cell(row=i, column=5, value=tar.estado.value.upper() if tar.estado else "")
        apply_badge(c_estado, tar.estado.value if tar.estado else "")
        style_row(ws_tar, i, 5, is_even=(i%2==0))

    # 5. Personal
    ws_pers = wb.create_sheet(title="Personal")
    start_row = draw_excel_dashboard_header(ws_pers, "Personal", perfil, kpi_data)
    
    headers_pers = ["NOMBRE / EMAIL", "ROL", "ESTADO"]
    for col, h in enumerate(headers_pers, 2):
        cell = ws_pers.cell(row=start_row, column=col, value=h)
        apply_table_header(cell)
        
    res_act = await db.execute(select(Usuario).where(Usuario.empresa_id == empresa_id))
    activos = res_act.scalars().all()
    
    row_idx = start_row + 1
    for u in activos:
        ws_pers.cell(row=row_idx, column=2, value=u.nombre)
        c_rol = ws_pers.cell(row=row_idx, column=3, value=u.rol.value.upper())
        apply_badge(c_rol, "ROLES")
        c_estado = ws_pers.cell(row=row_idx, column=4, value="ACTIVO")
        apply_badge(c_estado, "ACTIVO")
        style_row(ws_pers, row_idx, 4, is_even=(row_idx%2==0))
        row_idx += 1
        
    res_pend = await db.execute(select(Invitacion).where(Invitacion.empresa_id == empresa_id, Invitacion.estado == EstadoInvitacion.PENDIENTE))
    pendientes = res_pend.scalars().all()
    for p in pendientes:
        ws_pers.cell(row=row_idx, column=2, value=p.email)
        c_rol = ws_pers.cell(row=row_idx, column=3, value=(p.rol.value.upper() if p.rol else ""))
        apply_badge(c_rol, "ROLES")
        c_estado = ws_pers.cell(row=row_idx, column=4, value="PENDIENTE")
        apply_badge(c_estado, "PENDIENTE")
        style_row(ws_pers, row_idx, 4, is_even=(row_idx%2==0))
        row_idx += 1

    # Adjust widths for all sheets and FORCE font color
    for ws in wb.worksheets:
        ws.column_dimensions['A'].width = 3
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 40
        ws.column_dimensions['D'].width = 30
        ws.column_dimensions['E'].width = 25
        ws.column_dimensions['F'].width = 20
        ws.column_dimensions['G'].width = 15
        ws.column_dimensions['H'].width = 15
        ws.column_dimensions['I'].width = 15
        ws.column_dimensions['J'].width = 15
        
        # BRUTE FORCE FONT COLOR
        # Any row >= 11 is data
        for row in ws.iter_rows(min_row=11, max_row=ws.max_row, min_col=2, max_col=10):
            for cell in row:
                if cell.value is not None:
                    is_badge = False
                    if cell.fill and cell.fill.start_color and type(cell.fill.start_color.rgb) == str:
                        rgb_val = cell.fill.start_color.rgb
                        # Remove '00' alpha prefix if present
                        if len(rgb_val) == 8 and rgb_val.startswith("00"):
                            rgb_val = rgb_val[2:]
                        if rgb_val not in ["000000", "232B38", "2D3748", "1A202C"]:
                            is_badge = True
                    
                    if not is_badge:
                        if cell.column == 2: # Code
                            cell.font = Font(name="Consolas", size=10, color="94A3B8")
                        else:
                            cell.font = Font(name="Segoe UI", size=10, color="FFFFFF")

    import io
    from fastapi.responses import StreamingResponse
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    fecha_str = datetime.now().strftime("%Y%m%d_%H%M")
    empresa_str = perfil.nombre.replace(' ', '_') if perfil else 'Datos'
    filename = f"Exportacion_{empresa_str}_{fecha_str}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@router.get("/excel/presupuesto/{presupuesto_id}")
async def export_single_presupuesto(
    presupuesto_id: int,
    usuario: Usuario = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
    empresa_id: int = Depends(get_empresa_id)
):
    result = await db.execute(
        select(Presupuesto)
        .options(selectinload(Presupuesto.capitulos).selectinload(CapituloPresupuesto.partidas).selectinload(PartidaPresupuesto.lineas_medicion))
        .where(Presupuesto.id == presupuesto_id, Presupuesto.empresa_id == empresa_id)
    )
    presupuesto = result.scalar_one_or_none()
    if not presupuesto:
        raise HTTPException(status_code=404, detail="Presupuesto no encontrado o no autorizado")
        
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Presupuesto"
    ws.sheet_view.showGridLines = False
    
    font_title = Font(name="Calibri", size=16, bold=True, color="1F2937")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True, color="1F2937")
    font_normal = Font(name="Calibri", size=11, color="374151")
    font_medicion = Font(name="Calibri", size=10, color="6B7280", italic=True)
    
    fill_header = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    fill_capitulo = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    fill_subtotal = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
    
    border_bottom = Border(bottom=Side(style="thin", color="D1D5DB"))
    border_top = Border(top=Side(style="thin", color="9CA3AF"))
    
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 60
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 18
    
    ws.merge_cells("A1:G1")
    ws.cell(row=1, column=1, value=f"PRESUPUESTO: {presupuesto.nombre.upper()}").font = font_title
    
    current_row = 5
    
    headers = ["CÓDIGO", "DESCRIPCIÓN", "LÍNEAS", "UDS", "CANTIDAD", "PRECIO", "TOTAL"]
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=current_row, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        
    current_row += 1
    
    for capitulo in presupuesto.capitulos:
        ws.merge_cells(f"B{current_row}:G{current_row}")
        c_cod = ws.cell(row=current_row, column=1, value=str(capitulo.orden).zfill(2))
        c_cod.font = font_bold
        c_cod.fill = fill_capitulo
        c_desc = ws.cell(row=current_row, column=2, value=capitulo.nombre.upper())
        c_desc.font = font_bold
        c_desc.fill = fill_capitulo
        for c in range(1, 8):
            ws.cell(row=current_row, column=c).fill = fill_capitulo
        current_row += 1
        
        for partida in capitulo.partidas:
            c_cod = ws.cell(row=current_row, column=1, value=partida.codigo)
            c_cod.font = font_normal
            c_cod.border = border_bottom
            
            c_desc = ws.cell(row=current_row, column=2, value=partida.descripcion)
            c_desc.font = font_bold
            c_desc.border = border_bottom
            
            ws.cell(row=current_row, column=4, value=partida.unidad).font = font_normal
            ws.cell(row=current_row, column=4).border = border_bottom
            
            # Use cantidad_calculada safely!
            if hasattr(partida, "cantidad_calculada"):
                qty = partida.cantidad_calculada
            else:
                qty = partida.cantidad
            
            ws.cell(row=current_row, column=5, value=qty).font = font_normal
            ws.cell(row=current_row, column=5).border = border_bottom
            
            c_precio = ws.cell(row=current_row, column=6, value=partida.precio_unitario)
            c_precio.font = font_normal
            c_precio.border = border_bottom
            c_precio.number_format = '#,##0.00 €'
            
            # Need to recalculate importe because lazy load issues might still pop up if imported differently
            c_imp = ws.cell(row=current_row, column=7, value=float(qty) * float(partida.precio_con_descuento))
            c_imp.font = font_bold
            c_imp.border = border_bottom
            c_imp.number_format = '#,##0.00 €'
            
            # Format blank empty spaces
            for c in range(1, 8):
                ws.cell(row=current_row, column=c).border = border_bottom
            
            current_row += 1
            
            if partida.lineas_medicion:
                for lm in partida.lineas_medicion:
                    ws.cell(row=current_row, column=2, value=lm.comentario or "-").font = font_medicion
                    
                    detalle_parts = []
                    if lm.unidades is not None: detalle_parts.append(str(lm.unidades))
                    if lm.longitud is not None: detalle_parts.append(str(lm.longitud))
                    if lm.anchura is not None: detalle_parts.append(str(lm.anchura))
                    if lm.altura is not None: detalle_parts.append(str(lm.altura))
                    detalle = " x ".join(detalle_parts)
                    if not detalle: detalle = "-"
                    
                    ws.cell(row=current_row, column=3, value=detalle).font = font_medicion
                    ws.cell(row=current_row, column=5, value=lm.subtotal).font = font_medicion
                    current_row += 1
                    
        # Subtotal
        ws.merge_cells(f"A{current_row}:F{current_row}")
        c_sub = ws.cell(row=current_row, column=1, value=f"SUBTOTAL CAPÍTULO {str(capitulo.orden).zfill(2)}")
        c_sub.font = font_bold
        c_sub.alignment = Alignment(horizontal="right")
        
        # Calculate subtotal manually to avoid lazy loading
        total_cap = sum(float(getattr(p, "cantidad_calculada", getattr(p, "cantidad", 0))) * float(p.precio_con_descuento) for p in capitulo.partidas)
        
        c_val = ws.cell(row=current_row, column=7, value=total_cap)
        c_val.font = font_bold
        c_val.number_format = '#,##0.00 €'
        c_val.border = border_top
        for c in range(1, 8):
            ws.cell(row=current_row, column=c).fill = fill_subtotal
        current_row += 2

    import io
    from fastapi.responses import StreamingResponse
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    fecha_str = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"Presupuesto_{presupuesto.codigo or presupuesto.id}_{fecha_str}.xlsx"
    headers = {
        'Content-Disposition': f'attachment; filename="{filename}"'
    }
    return StreamingResponse(
        buffer, 
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 
        headers=headers
    )

