import os
import re
from typing import Any
import openpyxl
from pydantic import ValidationError
from app.services.importacion.models import (
    ResultadoAnalisis, MetadatosPresupuesto, CapituloIntermedio,
    PartidaIntermedia, ColumnMapping, LineaMedicionIntermedia
)

SEMANTIC_DICT = {
    "codigo": ["#", "código", "cod", "cod.", "ref", "referencia", "nº", "num", "n"],
    "descripcion": ["descripción", "descripcion", "desc", "concepto", "detalle", "partida"],
    "unidad": ["unidad", "ud", "ud.", "uds", "unit", "unit.", "u.m.", "medida"],
    "cantidad": ["cantidad", "cant", "cant.", "medición", "med.", "qty"],
    "precio_unitario": ["precio unitario", "precio", "p.u.", "precio unit", "p/u", "€/ud"],
    "descuento": ["desc. %", "descuento", "dto", "dto.", "dto %", "desc"],
    "importe": ["importe", "total", "subtotal", "total línea", "precio total", "amount"],
    "lineas": ["líneas", "lineas", "linea", "línea", "mediciones", "dimensiones"]
}

def clean_text(text: Any) -> str:
    if text is None:
        return ""
    return str(text).strip()

def parse_float(value: Any) -> float:
    if value is None or str(value).strip() == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    
    val_str = str(value).strip().lower()
    val_str = re.sub(r'[€%]', '', val_str).strip()
    
    if '.' in val_str and ',' in val_str:
        if val_str.rfind(',') > val_str.rfind('.'):
            val_str = val_str.replace('.', '').replace(',', '.')
        else:
            val_str = val_str.replace(',', '')
    elif ',' in val_str:
        parts = val_str.split(',')
        if len(parts) == 2 and len(parts[1]) != 3:
            val_str = val_str.replace(',', '.')
        elif len(parts) == 2 and len(parts[1]) == 3:
            val_str = val_str.replace(',', '.')
            
    try:
        return float(val_str)
    except ValueError:
        return 0.0

def parse_medicion_string(s: str) -> dict:
    # Example: '2.00 x 3.15 x 15.00 x 0.30'
    res = {"unidades": None, "longitud": None, "anchura": None, "altura": None}
    if not s:
        return res
    parts = re.split(r'[xX*]', str(s))
    nums = []
    for p in parts:
        val = parse_float(p)
        nums.append(val)
        
    if len(nums) > 0: res["unidades"] = nums[0]
    if len(nums) > 1: res["longitud"] = nums[1]
    if len(nums) > 2: res["anchura"] = nums[2]
    if len(nums) > 3: res["altura"] = nums[3]
    return res

async def analizar_excel(file_path: str) -> ResultadoAnalisis:
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        raise ValueError(f"Error al leer el archivo Excel: {str(e)}")
    
    if not wb.sheetnames:
        raise ValueError("El archivo Excel no contiene hojas.")
    
    ws = wb.active
    
    columnas_excel = []
    header_row_idx = -1
    column_mapping = ColumnMapping()
    
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        matches = 0
        current_mapping = {}
        for col_idx, cell_value in enumerate(row):
            if not cell_value:
                continue
            val_str = clean_text(cell_value).lower()
            for key, keywords in SEMANTIC_DICT.items():
                if val_str in keywords and key not in current_mapping:
                    current_mapping[key] = col_idx
                    matches += 1
                    break
        
        if matches >= 3:
            header_row_idx = row_idx
            # Add lineas mapping manually if found since ColumnMapping might not have it yet
            # Wait, column_mapping is a Pydantic model. We should use setattr. 
            # I'll just store the index in a variable since ColumnMapping doesn't officially support 'lineas'
            # Or I can dynamically add it, but it's easier to just track it locally
            lineas_col_idx = None
            for k, v in current_mapping.items():
                if k == 'lineas':
                    lineas_col_idx = v
                elif hasattr(column_mapping, k):
                    setattr(column_mapping, k, v)
            column_mapping_dict = current_mapping # Keep a full dictionary
            columnas_excel = [str(cell) for cell in row if cell is not None]
            break
            
    if header_row_idx == -1:
        raise ValueError("No se encontró la fila de cabeceras en el archivo Excel. Verifica el formato.")
        
    capitulos = []
    stack = []
    
    partidas_ok = 0
    partidas_warning = 0
    partidas_error = 0
    importe_total = 0.0
    
    iva = 21.0
    
    lineas_col_idx = column_mapping_dict.get('lineas')
    
    for row_idx, row in enumerate(ws.iter_rows(min_row=header_row_idx + 1, values_only=True), header_row_idx + 1):
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
            
        codigo = clean_text(row[column_mapping.codigo]) if column_mapping.codigo is not None and row[column_mapping.codigo] is not None else ""
        desc = clean_text(row[column_mapping.descripcion]) if column_mapping.descripcion is not None and row[column_mapping.descripcion] is not None else ""
        
        desc_lower = desc.lower()
        if "coste directo" in desc_lower or ("total" in desc_lower and "subtotal" not in desc_lower):
            continue
        if "iva" in desc_lower:
            m = re.search(r'(\d+(?:\.\d+)?)\s*%', desc_lower)
            if m:
                iva = float(m.group(1))
            continue
            
        unidad_raw = row[column_mapping.unidad] if column_mapping.unidad is not None else None
        cantidad_raw = row[column_mapping.cantidad] if column_mapping.cantidad is not None else None
        precio_raw = row[column_mapping.precio_unitario] if column_mapping.precio_unitario is not None else None
        desc_raw = row[column_mapping.descuento] if column_mapping.descuento is not None else None
        importe_raw = row[column_mapping.importe] if column_mapping.importe is not None else None
        lineas_raw = row[lineas_col_idx] if lineas_col_idx is not None else None
        
        # Determine if it's a chapter, partida, or measurement line
        is_chapter = False
        is_linea = False
        
        # If it has only description, it's a chapter
        if desc and not cantidad_raw and not precio_raw and not importe_raw and not lineas_raw:
            is_chapter = True
            
        # If it has NO price, NO importe, NO code, but has quantity, it's a linea de medicion
        # OR if it explicitly has something in the 'LÍNEAS' column
        if (not codigo and (not precio_raw or parse_float(precio_raw) == 0.0) and (not importe_raw or parse_float(importe_raw) == 0.0)) and (cantidad_raw is not None or lineas_raw is not None):
            # Only if we have a parent partida to attach it to
            if stack and stack[-1].partidas:
                is_linea = True
                is_chapter = False
            
        if is_linea:
            # Parse the measurement line
            comentario = desc if desc != "-" else ""
            cantidad_val = parse_float(cantidad_raw)
            dims = parse_medicion_string(clean_text(lineas_raw))
            
            linea_medicion = LineaMedicionIntermedia(
                comentario=comentario,
                unidades=dims['unidades'],
                longitud=dims['longitud'],
                anchura=dims['anchura'],
                altura=dims['altura'],
                subtotal=cantidad_val
            )
            stack[-1].partidas[-1].lineas_medicion.append(linea_medicion)
            continue
            
        if is_chapter:
            level = 1
            if codigo:
                parts = [p for p in re.split(r'[.-]', codigo) if p.strip()]
                if len(parts) > 0:
                    level = len(parts)
                    
            nuevo_cap = CapituloIntermedio(
                codigo=codigo,
                nombre=desc,
                orden=0,
                partidas=[],
                subcapitulos=[]
            )
            setattr(nuevo_cap, '_level', level)
            
            while stack and getattr(stack[-1], '_level', 1) >= level:
                stack.pop()
                
            if not stack:
                nuevo_cap.orden = len(capitulos) + 1
                capitulos.append(nuevo_cap)
            else:
                parent = stack[-1]
                nuevo_cap.orden = len(parent.subcapitulos) + 1
                parent.subcapitulos.append(nuevo_cap)
                
            stack.append(nuevo_cap)
            
        elif desc or codigo:
            if not stack:
                dummy = CapituloIntermedio(
                    codigo="01",
                    nombre="Capítulo General",
                    orden=len(capitulos) + 1,
                    partidas=[],
                    subcapitulos=[]
                )
                setattr(dummy, '_level', 1)
                capitulos.append(dummy)
                stack.append(dummy)
                
            cantidad = parse_float(cantidad_raw)
            precio = parse_float(precio_raw)
            descuento = parse_float(desc_raw)
            importe = parse_float(importe_raw)
            unidad = clean_text(unidad_raw) or "ud"
            
            warnings = []
            status = "ok"
            
            calculated_importe = cantidad * precio * (1 - descuento/100)
            if importe == 0.0 and calculated_importe > 0:
                importe = calculated_importe
                
            if importe > 0:
                diff = abs(calculated_importe - importe)
                if diff > 0.01 * importe:
                    warnings.append(f"El importe ({importe:.2f}) no coincide con cantidad x precio x (1 - dto) ({calculated_importe:.2f})")
                    status = "warning"
                    partidas_warning += 1
                else:
                    partidas_ok += 1
            else:
                if not (precio == 0.0 and importe == 0.0): # Valid for blind budgets
                    status = "error"
                    warnings.append("Falta cantidad o precio unitario, o su producto es cero")
                    partidas_error += 1
                
            partida = PartidaIntermedia(
                codigo=codigo,
                descripcion=desc,
                unidad=unidad,
                cantidad=cantidad,
                precio_unitario=precio,
                descuento_porcentaje=descuento,
                importe=importe,
                status=status,
                warnings=warnings,
                fila_origen=row_idx,
                lineas_medicion=[]
            )
            
            stack[-1].partidas.append(partida)
            importe_total += importe
        
    metadatos = MetadatosPresupuesto(
        nombre_obra=os.path.basename(file_path),
        iva=iva
    )
    
    total_partidas = sum(len(c.partidas) for c in capitulos)
    
    return ResultadoAnalisis(
        metadatos=metadatos,
        capitulos=capitulos,
        column_mapping=column_mapping,
        total_capitulos=len(capitulos),
        total_partidas=total_partidas,
        partidas_ok=partidas_ok,
        partidas_warning=partidas_warning,
        partidas_error=partidas_error,
        importe_total_detectado=importe_total,
        warnings_globales=[],
        columnas_excel=columnas_excel
    )

