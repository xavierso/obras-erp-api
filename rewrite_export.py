import re

NEW_CODE = '''def draw_excel_dashboard_header(ws, title: str, perfil: Empresa, kpi_data: dict):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    
    BG_MAIN = "1A202C"      # Fondo muy oscuro
    BG_CARD = "2D3748"      # Gris oscuro para tarjetas
    ACCENT = "0D9488"       # Teal / Cyan para el acento
    TEXT_MAIN = "F8FAFC"
    TEXT_MUTED = "94A3B8"
    
    font_family = "Segoe UI"
    
    fill_bg = PatternFill(start_color=BG_MAIN, end_color=BG_MAIN, fill_type="solid")
    fill_card = PatternFill(start_color=BG_CARD, end_color=BG_CARD, fill_type="solid")
    
    ws.sheet_view.showGridLines = False
    
    # Fondo global
    for r in range(1, 150):
        for c in range(1, 15):
            ws.cell(row=r, column=c).fill = fill_bg

    # --- ROW 1: Título Principal ---
    ws.cell(row=1, column=2, value="DIAM GESTIÓN").font = Font(name=font_family, size=22, bold=True, color=TEXT_MAIN)
    ws.cell(row=1, column=4, value="DE OBRAS").font = Font(name=font_family, size=22, bold=True, color=ACCENT)
    
    # --- ROW 3: Logo / Proyecto y Tabs ---
    ws.cell(row=3, column=2, value="Empresa").font = Font(name=font_family, size=9, color=TEXT_MUTED)
    ws.cell(row=4, column=2, value=perfil.nombre if perfil else "DIAM ERP").font = Font(name=font_family, size=11, color=TEXT_MAIN)
    
    tabs = ["Presupuestos", "Obras", "Incidencias", "Tareas", "Personal"]
    col_idx = 4
    for tab in tabs:
        c = ws.cell(row=3, column=col_idx, value=tab.upper())
        c.font = Font(name=font_family, size=11, color=TEXT_MAIN if tab == title else TEXT_MUTED)
        if tab == title:
            c.fill = PatternFill(start_color=ACCENT, end_color=ACCENT, fill_type="solid")
            c.font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=3, start_column=col_idx, end_row=4, end_column=col_idx)
        col_idx += 1
        
    # --- ROW 6-8: Tarjetas KPI ---
    # Box 1: Obras
    for r in range(6, 9):
        for c in range(2, 4):
            ws.cell(row=r, column=c).fill = fill_card
    ws.cell(row=6, column=2, value="OBRAS ACTIVAS").font = Font(name=font_family, size=10, color=TEXT_MUTED)
    ws.cell(row=7, column=2, value=kpi_data.get('obras_activas', 0)).font = Font(name=font_family, size=20, bold=True, color=ACCENT)
    ws.cell(row=7, column=2).alignment = Alignment(vertical="center")
    
    # Box 2: Presupuesto
    for r in range(6, 9):
        for c in range(4, 6):
            ws.cell(row=r, column=c).fill = fill_card
    ws.cell(row=6, column=4, value="PRESUPUESTADO TOTAL").font = Font(name=font_family, size=10, color=TEXT_MUTED)
    
    val_total = kpi_data.get('presupuesto_total', 0)
    c_monto = ws.cell(row=7, column=4, value=val_total)
    c_monto.font = Font(name=font_family, size=20, bold=True, color=TEXT_MAIN)
    c_monto.number_format = '#,##0.00 €'
    c_monto.alignment = Alignment(vertical="center")
    
    # Box 3: Incidencias
    for r in range(6, 9):
        for c in range(6, 8):
            ws.cell(row=r, column=c).fill = fill_card
    ws.cell(row=6, column=6, value="INCIDENCIAS ABIERTAS").font = Font(name=font_family, size=10, color=TEXT_MUTED)
    ws.cell(row=7, column=6, value=kpi_data.get('incidencias_abiertas', 0)).font = Font(name=font_family, size=18, bold=True, color="EF4444")
    ws.cell(row=7, column=6).alignment = Alignment(vertical="center")
    
    # Box 4: Equipo
    for r in range(6, 9):
        for c in range(8, 10):
            ws.cell(row=r, column=c).fill = fill_card
    ws.cell(row=6, column=8, value="EQUIPO EN CAMPO").font = Font(name=font_family, size=10, color=TEXT_MUTED)
    ws.cell(row=7, column=8, value=kpi_data.get('equipo', 0)).font = Font(name=font_family, size=18, bold=True, color="10B981")
    ws.cell(row=7, column=8).alignment = Alignment(vertical="center")
    
    # --- ROW 10: Barra de Título de la tabla ---
    ws.cell(row=10, column=2, value=f"RESUMEN DE {title.upper()}").font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    teal_fill = PatternFill(start_color=ACCENT, end_color=ACCENT, fill_type="solid")
    for c in range(2, 10):
        ws.cell(row=10, column=c).fill = teal_fill

    return 11

def apply_table_header(cell):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    cell.fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid") # Teal header
    cell.font = Font(color="FFFFFF", bold=True, name="Segoe UI", size=11)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.border = Border(bottom=Side(style="medium", color="0F766E"))

def apply_badge(cell, estado: str):
    from openpyxl.styles import Font, PatternFill, Alignment
    estado = (estado or "").upper()
    font_family = "Segoe UI"
    
    if estado in ["EN_EJECUCION", "ACTIVO"]:
        cell.fill = PatternFill(start_color="0EA5E9", end_color="0EA5E9", fill_type="solid") 
        cell.font = Font(name=font_family, color="FFFFFF", bold=True, size=9)
    elif estado in ["PENDIENTE", "PENDIENTE_APROBACION", "BORRADOR", "EN_ESTUDIO", "ENVIADO"]:
        cell.fill = PatternFill(start_color="EAB308", end_color="EAB308", fill_type="solid") 
        cell.font = Font(name=font_family, color="000000", bold=True, size=9)
    elif estado in ["INCIDENCIA", "BLOQUEADO", "CANCELADO"]:
        cell.fill = PatternFill(start_color="EF4444", end_color="EF4444", fill_type="solid") 
        cell.font = Font(name=font_family, color="FFFFFF", bold=True, size=9)
    elif estado in ["COMPLETADO", "FINALIZADO", "APROBADO", "RESUELTA"]:
        cell.fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid") 
        cell.font = Font(name=font_family, color="FFFFFF", bold=True, size=9)
    else:
        cell.fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid") 
        cell.font = Font(name=font_family, color="FFFFFF", bold=True, size=9)
        
    cell.alignment = Alignment(horizontal="center", vertical="center")

@router.get("/excel")
async def export_to_excel(
    admin: Usuario = Depends(require_director),
    db: AsyncSession = Depends(get_db),
    empresa_id: int = Depends(get_empresa_id)
):
    result = await db.execute(select(Empresa).where(Empresa.id == empresa_id))
    perfil = result.scalar_one_or_none()

    # Pre-calculate KPIs
    res_obras = await db.execute(select(Obra).where(Obra.empresa_id == empresa_id))
    obras = res_obras.scalars().all()
    obras_activas = sum(1 for o in obras if o.estado and o.estado.value not in ["cancelado", "finalizado"])
    
    res_pres = await db.execute(
        select(Presupuesto).options(selectinload(Presupuesto.capitulos).selectinload(CapituloPresupuesto.partidas)).where(Presupuesto.empresa_id == empresa_id)
    )
    presupuestos = res_pres.scalars().all()
    presupuesto_total = sum(p.total or 0 for p in presupuestos)
    
    res_inc = await db.execute(select(Incidencia).options(selectinload(Incidencia.obra)).where(Incidencia.empresa_id == empresa_id))
    incidencias = res_inc.scalars().all()
    incidencias_abiertas = sum(1 for i in incidencias if i.estado and i.estado.value not in ["resuelta"])
    
    res_usu = await db.execute(select(Usuario).where(Usuario.empresa_id == empresa_id))
    equipo_count = len(res_usu.scalars().all())
    
    kpi_data = {
        "obras_activas": obras_activas,
        "presupuesto_total": presupuesto_total,
        "incidencias_abiertas": incidencias_abiertas,
        "equipo": equipo_count
    }

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    BG_MAIN = "1A202C"
    BG_CARD = "2D3748"
    TEXT_MAIN = "F8FAFC"
    
    def style_row(ws, row_idx, max_col, is_even):
        bg = BG_CARD if is_even else "232B38"
        from openpyxl.styles import PatternFill, Font, Border, Side
        fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
        for c in range(2, max_col + 1):
            cell = ws.cell(row=row_idx, column=c)
            # if no explicit fill like badge
            if not cell.fill.start_color or type(cell.fill.start_color.rgb) != str or cell.fill.start_color.rgb == "00000000":
                cell.fill = fill
            if not cell.font or not cell.font.color or cell.font.color.rgb == "00000000":
                cell.font = Font(name="Segoe UI", size=10, color=TEXT_MAIN)
            cell.border = Border(bottom=Side(style="thin", color="1A202C"), right=Side(style="thin", color="1A202C"))

    # 1. Presupuestos
    ws_pres = wb.create_sheet(title="Presupuestos")
    start_row = draw_excel_dashboard_header(ws_pres, "Presupuestos", perfil, kpi_data)
    
    headers = ["CÓDIGO", "NOMBRE", "ESTADO", "MONTO"]
    for col, h in enumerate(headers, 2):
        cell = ws_pres.cell(row=start_row, column=col, value=h)
        apply_table_header(cell)
        
    for i, p in enumerate(presupuestos, start=start_row + 1):
        ws_pres.cell(row=i, column=2, value=p.codigo or "").font = Font(name="Consolas", color="94A3B8")
        ws_pres.cell(row=i, column=3, value=p.nombre)
        
        c_estado = ws_pres.cell(row=i, column=4, value=p.estado.value.upper() if p.estado else "")
        apply_badge(c_estado, p.estado.value if p.estado else "")
        
        c_monto = ws_pres.cell(row=i, column=5, value=p.total)
        c_monto.number_format = '#,##0.00 €'
        c_monto.alignment = Alignment(horizontal="right")
        
        style_row(ws_pres, i, 5, is_even=(i%2==0))
        
    # 2. Obras
    ws_obras = wb.create_sheet(title="Obras")
    start_row = draw_excel_dashboard_header(ws_obras, "Obras", perfil, kpi_data)
    
    headers_obras = ["CÓDIGO", "NOMBRE", "CLIENTE", "ESTADO"]
    for col, h in enumerate(headers_obras, 2):
        cell = ws_obras.cell(row=start_row, column=col, value=h)
        apply_table_header(cell)
        
    for i, o in enumerate(obras, start=start_row + 1):
        ws_obras.cell(row=i, column=2, value=o.codigo).font = Font(name="Consolas", color="94A3B8")
        ws_obras.cell(row=i, column=3, value=o.nombre)
        ws_obras.cell(row=i, column=4, value=o.cliente or "")
        c_estado = ws_obras.cell(row=i, column=5, value=o.estado.value.upper() if o.estado else "")
        apply_badge(c_estado, o.estado.value if o.estado else "")
        style_row(ws_obras, i, 5, is_even=(i%2==0))

    # 3. Incidencias
    ws_inc = wb.create_sheet(title="Incidencias")
    start_row = draw_excel_dashboard_header(ws_inc, "Incidencias", perfil, kpi_data)
    
    headers_inc = ["CÓDIGO", "TÍTULO", "OBRA", "ESTADO"]
    for col, h in enumerate(headers_inc, 2):
        cell = ws_inc.cell(row=start_row, column=col, value=h)
        apply_table_header(cell)
        
    for i, inc in enumerate(incidencias, start=start_row + 1):
        ws_inc.cell(row=i, column=2, value=inc.codigo).font = Font(name="Consolas", color="94A3B8")
        ws_inc.cell(row=i, column=3, value=inc.titulo)
        ws_inc.cell(row=i, column=4, value=inc.obra.nombre if inc.obra else "")
        c_estado = ws_inc.cell(row=i, column=5, value=inc.estado.value.upper() if inc.estado else "")
        apply_badge(c_estado, inc.estado.value if inc.estado else "")
        style_row(ws_inc, i, 5, is_even=(i%2==0))

    # 4. Tareas
    ws_tar = wb.create_sheet(title="Tareas")
    start_row = draw_excel_dashboard_header(ws_tar, "Tareas", perfil, kpi_data)
    
    res_tar = await db.execute(select(Tarea).options(selectinload(Tarea.obra)).where(Tarea.empresa_id == empresa_id))
    tareas = res_tar.scalars().all()
    
    headers_tar = ["TÍTULO", "OBRA", "FECHA LÍMITE", "ESTADO"]
    for col, h in enumerate(headers_tar, 2):
        cell = ws_tar.cell(row=start_row, column=col, value=h)
        apply_table_header(cell)
        
    for i, tar in enumerate(tareas, start=start_row + 1):
        ws_tar.cell(row=i, column=2, value=tar.titulo)
        ws_tar.cell(row=i, column=3, value=tar.obra.nombre if tar.obra else "")
        ws_tar.cell(row=i, column=4, value=tar.fecha_limite.strftime("%Y-%m-%d") if tar.fecha_limite else "")
        c_estado = ws_tar.cell(row=i, column=5, value=tar.estado.value.upper() if tar.estado else "")
        apply_badge(c_estado, tar.estado.value if tar.estado else "")
        style_row(ws_tar, i, 5, is_even=(i%2==0))

    # 5. Personal
    ws_pers = wb.create_sheet(title="Personal")
    start_row = draw_excel_dashboard_header(ws_pers, "Personal", perfil, kpi_data)
    
    headers_pers = ["NOMBRE / EMAIL", "ROL", "ESTADO"]
    for col, h in enumerate(headers_pers, 2):
        cell = ws_pers.cell(row=start_row, column=col, value=h)
        apply_table_header(cell)
        
    res_act = await db.execute(select(Usuario).where(Usuario.empresa_id == empresa_id))
    activos = res_act.scalars().all()
    
    row_idx = start_row + 1
    for u in activos:
        ws_pers.cell(row=row_idx, column=2, value=u.nombre)
        c_rol = ws_pers.cell(row=row_idx, column=3, value=u.rol.value.upper())
        apply_badge(c_rol, "ROLES")
        c_estado = ws_pers.cell(row=row_idx, column=4, value="ACTIVO")
        apply_badge(c_estado, "ACTIVO")
        style_row(ws_pers, row_idx, 4, is_even=(row_idx%2==0))
        row_idx += 1
        
    res_pend = await db.execute(select(Invitacion).where(Invitacion.empresa_id == empresa_id, Invitacion.estado == EstadoInvitacion.PENDIENTE))
    pendientes = res_pend.scalars().all()
    for p in pendientes:
        ws_pers.cell(row=row_idx, column=2, value=p.email)
        c_rol = ws_pers.cell(row=row_idx, column=3, value=(p.rol.value.upper() if p.rol else ""))
        apply_badge(c_rol, "ROLES")
        c_estado = ws_pers.cell(row=row_idx, column=4, value="PENDIENTE")
        apply_badge(c_estado, "PENDIENTE")
        style_row(ws_pers, row_idx, 4, is_even=(row_idx%2==0))
        row_idx += 1

    # Adjust widths for all sheets
    for ws in wb.worksheets:
        ws.column_dimensions['A'].width = 3
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 40
        ws.column_dimensions['D'].width = 30
        ws.column_dimensions['E'].width = 25
        ws.column_dimensions['F'].width = 20
        ws.column_dimensions['G'].width = 15
        ws.column_dimensions['H'].width = 15
        ws.column_dimensions['I'].width = 15
        ws.column_dimensions['J'].width = 15

    import io
    from fastapi.responses import StreamingResponse
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    fecha_str = datetime.now().strftime("%Y%m%d_%H%M")
    empresa_str = perfil.nombre.replace(' ', '_') if perfil else 'Datos'
    filename = f"Exportacion_{empresa_str}_{fecha_str}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
'''

with open("app/routers/exportacion.py", "r", encoding="utf-8") as f:
    content = f.read()

start_idx = content.find("def draw_excel_dashboard_header")
end_idx = content.find("@router.get(\"/excel/presupuesto/{presupuesto_id}\")")

if start_idx != -1 and end_idx != -1:
    new_content = content[:start_idx] + NEW_CODE + "\n" + content[end_idx:]
    with open("app/routers/exportacion.py", "w", encoding="utf-8") as f:
        f.write(new_content)
    print("Success")
else:
    print("Could not find boundaries")
