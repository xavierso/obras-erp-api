import openpyxl
from openpyxl.styles import Font, PatternFill

wb = openpyxl.Workbook()
ws = wb.active

ws.cell(1, 1, "Test default")
c2 = ws.cell(2, 1, "Test white")
c2.font = Font(color="FFFFFF")
c2.fill = PatternFill(start_color="2D3748", end_color="2D3748", fill_type="solid")

wb.save("test_font.xlsx")
print("Saved")
